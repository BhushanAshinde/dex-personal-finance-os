import csv
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..database import get_db
from ..deps import current_user
from ..models import Category, Transaction, User, Wallet
from ..services.finance import get_or_create_category

router = APIRouter(prefix="/statements", tags=["statements"])

MAX_IMPORT_BYTES = 10 * 1024 * 1024


def get_transactions(user_id: int, db: Session, wallet_id: int | None, date_from: date | None, date_to: date | None, transaction_type: str | None):
    query = select(Transaction, Wallet).join(Wallet, Wallet.id == Transaction.wallet_id).where(Transaction.user_id == user_id)
    if wallet_id is not None:
        query = query.where(Transaction.wallet_id == wallet_id)
    if date_from:
        query = query.where(Transaction.transaction_date >= datetime.combine(date_from, time.min))
    if date_to:
        query = query.where(Transaction.transaction_date <= datetime.combine(date_to, time.max))
    if transaction_type:
        query = query.where(Transaction.kind == transaction_type.lower())
    return db.execute(query.order_by(Transaction.transaction_date, Transaction.id)).all()


def row_values(rows):
    running = {}
    result = []
    for transaction, wallet in rows:
        running.setdefault(wallet.id, wallet.opening_balance or 0)
        debit = transaction.amount if transaction.kind == "expense" or transaction.transfer_direction == "out" else 0
        credit = transaction.amount if transaction.kind in {"income", "refund"} or transaction.transfer_direction == "in" else 0
        running[wallet.id] += credit - debit
        result.append({
            "transaction_id": transaction.id,
            "date": transaction.transaction_date.strftime("%Y-%m-%d"),
            "wallet": wallet.name,
            "bank": wallet.bank_name or "",
            "account_last4": wallet.account_last4 or "",
            "type": transaction.kind,
            "category": transaction.category.name if transaction.category else "",
            "merchant": transaction.merchant or "",
            "description": transaction.note or "",
            "debit": debit,
            "credit": credit,
            "balance": running[wallet.id],
            "currency": wallet.currency,
        })
    return result


def validate_wallet(wallet_id: int | None, user_id: int, db: Session):
    if wallet_id is not None and not db.scalar(select(Wallet).where(Wallet.id == wallet_id, Wallet.user_id == user_id)):
        raise HTTPException(404, "Wallet not found")


def parse_import_rows(filename: str, content: bytes):
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(413, "Statement file is too large")
    extension = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if extension == "csv":
        return list(csv.DictReader(StringIO(content.decode("utf-8-sig"))))
    if extension in {"xlsx", "xlsm"}:
        sheet = load_workbook(BytesIO(content), read_only=True, data_only=True).active
        rows = list(sheet.values)
        if not rows:
            return []
        headers = [str(value or "").strip().lower() for value in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:]]
    raise HTTPException(415, "Only CSV and XLSX files are supported")


def row_value(row, *names):
    normalized = {str(key).strip().lower(): value for key, value in row.items()}
    for name in names:
        if normalized.get(name) not in (None, ""):
            return normalized[name]
    return None


@router.post("/import")
def import_statement(wallet_id: int = Form(...), file: UploadFile = File(...), user: User = Depends(current_user), db: Session = Depends(get_db)):
    wallet = db.scalar(select(Wallet).where(Wallet.id == wallet_id, Wallet.user_id == user.id, Wallet.active == True))
    if not wallet:
        raise HTTPException(404, "Active wallet not found")
    rows = parse_import_rows(file.filename or "", file.file.read())
    existing = {(tx.transaction_date.date(), str(tx.amount), tx.note or "", tx.reference_number or "") for tx in db.scalars(select(Transaction).where(Transaction.user_id == user.id, Transaction.wallet_id == wallet.id)).all()}
    imported = duplicates = invalid = 0
    for row in rows:
        try:
            raw_date = row_value(row, "date", "transaction date", "transaction_date")
            transaction_date = raw_date if isinstance(raw_date, datetime) else datetime.fromisoformat(str(raw_date).replace("/", "-")[:10])
            amount = Decimal(str(row_value(row, "amount", "debit", "credit")).replace(",", ""))
            description = str(row_value(row, "description", "narration", "details", "remarks") or "").strip()
            reference = str(row_value(row, "reference", "reference number", "ref") or "").strip() or None
            kind = "income" if row_value(row, "credit", "income") not in (None, "", 0, "0") and row_value(row, "debit") in (None, "", 0, "0") else "expense"
            signature = (transaction_date.date(), str(abs(amount)), description, reference or "")
            if signature in existing:
                duplicates += 1
                continue
            category = get_or_create_category(db, str(row_value(row, "category") or "Other"))
            amount = abs(amount)
            db.add(Transaction(user_id=user.id, wallet_id=wallet.id, category_id=category.id, kind=kind, amount=amount, merchant=row_value(row, "merchant"), note=description, reference_number=reference, transaction_date=transaction_date))
            wallet.balance += amount if kind == "income" else -amount
            existing.add(signature)
            imported += 1
        except (TypeError, ValueError, InvalidOperation):
            invalid += 1
    db.commit()
    return {"total_rows": len(rows), "imported": imported, "duplicates": duplicates, "invalid": invalid}


@router.get("/export/csv")
def export_csv(wallet_id: int | None = None, date_from: date | None = None, date_to: date | None = None, transaction_type: str | None = None, user: User = Depends(current_user), db: Session = Depends(get_db)):
    validate_wallet(wallet_id, user.id, db)
    output = StringIO()
    columns = ["transaction_id", "date", "wallet", "bank", "account_last4", "type", "category", "merchant", "description", "debit", "credit", "balance", "currency"]
    writer = csv.DictWriter(output, fieldnames=columns)
    writer.writeheader()
    writer.writerows(row_values(get_transactions(user.id, db, wallet_id, date_from, date_to, transaction_type)))
    return Response(output.getvalue(), media_type="text/csv", headers={"Content-Disposition": "attachment; filename=dex-statement.csv"})


@router.get("/export/excel")
def export_excel(wallet_id: int | None = None, date_from: date | None = None, date_to: date | None = None, transaction_type: str | None = None, user: User = Depends(current_user), db: Session = Depends(get_db)):
    validate_wallet(wallet_id, user.id, db)
    rows = row_values(get_transactions(user.id, db, wallet_id, date_from, date_to, transaction_type))
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Statement Summary"
    summary.append(["DEX PERSONAL FINANCE", ""])
    summary.append(["Period", f"{date_from or 'All time'} to {date_to or 'Today'}"])
    summary.append(["Transactions", len(rows)])
    summary.append(["Total Income", sum(row["credit"] for row in rows if row["type"] in {"income", "refund"})])
    summary.append(["Total Expense", sum(row["debit"] for row in rows if row["type"] == "expense")])
    summary.append(["Closing Balance", rows[-1]["balance"] if rows else 0])
    summary["A1"].font = Font(bold=True, size=16)
    sheet = workbook.create_sheet("Transactions")
    columns = list(rows[0].keys()) if rows else ["transaction_id", "date", "wallet", "bank", "account_last4", "type", "category", "merchant", "description", "debit", "credit", "balance", "currency"]
    sheet.append(columns)
    for cell in sheet[1]: cell.font = Font(bold=True)
    for row in rows: sheet.append([row[column] for column in columns])
    sheet.auto_filter.ref = sheet.dimensions
    buffer = BytesIO()
    workbook.save(buffer)
    return Response(buffer.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=dex-statement.xlsx"})


@router.get("/export/pdf")
def export_pdf(wallet_id: int | None = None, date_from: date | None = None, date_to: date | None = None, transaction_type: str | None = None, user: User = Depends(current_user), db: Session = Depends(get_db)):
    validate_wallet(wallet_id, user.id, db)
    rows = row_values(get_transactions(user.id, db, wallet_id, date_from, date_to, transaction_type))
    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=24, leftMargin=24, topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    content = [Paragraph("DEX PERSONAL FINANCE", styles["Title"]), Paragraph(f"Account Statement | {date_from or 'All time'} to {date_to or 'Today'}", styles["Normal"]), Spacer(1, 12)]
    headers = ["Date", "Transaction", "Wallet", "Type", "Category", "Debit", "Credit", "Balance"]
    data = [headers] + [[row["date"], row["transaction_id"], row["wallet"], row["type"], row["category"], row["debit"], row["credit"], row["balance"]] for row in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1c2521")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 8)]))
    content.append(table)
    document.build(content)
    return Response(buffer.getvalue(), media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=dex-statement.pdf"})
